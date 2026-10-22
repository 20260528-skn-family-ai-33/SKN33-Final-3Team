"""상위어·하위어 후보 쌍을 규칙으로 만들고 관계를 판정합니다. build_hypernyms.py가 이 결과로 상위어 사전을 만듭니다.
확인용 표: data/news_keywords/review/hierarchy_review.xlsx (검수 파일은 읽지 않음)

대상 쌍 (모두 규칙, 저장소 밖 파일 없이 다시 만들 수 있음)
  ① 동의어 사전(build_synonyms.py 결과)의 쌍 — 접미사 규칙 줄은 바로 위 기본 단어와의 쌍도
  ② 임베딩 이웃: 기사 10건 이상 주제어를 OpenAI text-embedding-3-large로 임베딩, 말마다 가장 가까운 3개 → 유사도 상위 232쌍
  ③ 동시 등장: 넓은 말(rule_data.BROAD_WORDS)을 빼고, 둘 다 기사 5건 이상 · 함께 5건 이상 · 적은 쪽 기사의 50% 이상 같이 나옴
     → 유사도 상위 71쌍
  (232·71은 예전 후보(e5 유사도 기준)와 같은 개수. 임베딩은 data/news_keywords/word_vectors.npz에 저장해 다시 쓰고 새 말만 임베딩)
관계 판정 순서 (코드 규칙, 위가 우선)
  1. 반대·짝      반대말 쌍(rule_data.SYNONYM_OPPOSITES: 수출↔수입, 구직↔구인 …) — 상하위 아님
  2. 글자 포함    짧은 말이 긴 말 안에 그대로 있으면 짧은 말 = 상위어
     · 범위(산업·시장·업)    뒤에 산업·시장·업만 붙음 (화장품 → 화장품시장)            — 지금 접미사 규칙
     · 같은 대상(범용 접미사) 뒤에 성·품·률·율·회·케어만 붙음 (의약 → 의약품)         — 상하위 아님, 같은 수준
     · 종류(앞에 수식어)      앞에만 붙음 (화장품 → 기초화장품, 자전거 → 전기자전거)    — 접두사 규칙
     · 세부 주제(뒤에 붙음)   뒤에 다른 말이 붙음 (전기차 → 전기차보조금)
     · 세부 주제(앞뒤 붙음)   앞뒤 모두 붙음
  3. 같은 뜻 묶음  build_synonyms.GROUPS(괄호·표기·LLM 판단)로 묶인 쌍 — 같은 수준 (EV = 전기차처럼 약어는 동시 등장이 한쪽으로 쏠려도 동의어)
  4. 동시 등장     글자가 겹치지 않으면 기사 동시 등장의 쏠림으로 방향을 봄 (kw_clean, 동의어 적용 전 표기)
                  P(넓은 말 | 좁은 말) ≥ 0.3 이고 P(좁은 말 | 넓은 말) ≤ 0.1, 넓은 말 기사가 3배 이상 → 좁은 말 = 하위어 (반려견 → 반려동물)
  5. 그 밖        같은 수준(판단 보류)
  6. 상위어가 넓은 말(BROAD_WORDS: 수출·투자·기술 …)이거나 한 글자(물·차 …)면 제외
  마지막에 팀 지정(코드 안 예외: MANUAL_CHILDREN · MANUAL_SEPARATE)이 덮어씀

실행: python backend/app/services/news_keywords/build/hierarchy_candidates.py [--output 경로]
"""
import argparse
import importlib.util
from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
NK = HERE.parent
ROOT = Path(__file__).resolve().parents[5]
DATA = ROOT / "data" / "news_keywords"
OUT = DATA / "review" / "hierarchy_review.xlsx"
VECTORS = DATA / "word_vectors.npz"   # 주제어 임베딩 저장 (git 제외)
ENV_PATH = ROOT / ".env"
EMBED_MODEL = "text-embedding-3-large"


def _load(name):
    spec = importlib.util.spec_from_file_location(name, HERE / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


bs = _load("build_synonyms")
xu = _load("xlsx_util")
rd = bs.rd

RANGE_SUFFIX = ("산업", "시장", "업")
SAME_SUFFIX = ("성", "품", "률", "율", "회", "케어")
HIER = {"범위(산업·시장·업)", "종류(앞에 수식어)", "세부 주제(뒤에 붙음)", "세부 주제(앞뒤 붙음)", "포함(동시 등장)", "팀 지정"}
EXCLUDED = "제외(넓은·한 글자 상위어)"
EMB_MIN, EMB_TOP, EMB_PAIRS = 10, 3, 232                      # 임베딩 이웃: 기사 10건 이상 · 말마다 3개 · 상위 232쌍
CO_MIN_DOC, CO_MIN_BOTH, CO_SHARE, CO_PAIRS = 5, 5, 0.5, 71   # 동시 등장 후보

# 팀 지정 (2026-10-21): 하위어를 이 상위어 아래로. 규칙 결과보다 우선하고, 하위어의 다른 상위어 연결은 뺌
MANUAL_CHILDREN = {
    "미용": ["미용기기", "미용의료기기"],
    "온라인쇼핑": ["온라인쇼핑시장", "온라인시장"],
    "의료": ["의료시장", "의료관광", "의료장비", "의료기기산업", "의료미용시장", "의료장비시장"],
    "의약품": ["의약", "의약산업", "의약품시장", "의약품산업"],
    "전기차": ["전기차인프라", "전기차충전"],
    "카페": ["카페시장", "커피머신", "커피머신시장"],
    "반려동물": ["펫용품", "펫용품시장", "펫케어", "펫케어시장"],
}
# 예외: 규칙 후보에는 안 나오지만 팀이 적용한 기존 상위어 사전(2026-10-21)에 있던 연결. 넓은 말·한 글자 상위어는 뺌
KEPT_CHILDREN = {
    'ESS': ['BESS'],
    'PB': ['PB상품'],
    'Z세대': ['MZ세대'],
    '건강': ['건강기능식품'],
    '건설': ['건축자재'],
    '고용': ['근로자', '노동법'],
    '공급망': ['공급망위기', '글로벌공급망'],
    '교통': ['대중교통'],
    '금리': ['기준금리'],
    '농업': ['농산물'],
    '리튬': ['수산화리튬', '탄산리튬'],
    '모빌리티': ['마이크로모빌리티'],
    '반도체': ['반도체공급망', '반도체소재'],
    '방산': ['방산전시회'],
    '뷰티': ['뷰티디바이스'],
    '비만': ['비만치료제'],
    '수소': ['수소차'],
    '식품': ['농식품', '식료품'],
    '영화': ['영화관'],
    '예산': ['정부예산'],
    '의료': ['보건', '의료기기', '의료기술', '의료바이오', '진단기기'],
    '의약품': ['일반의약품', '제약'],
    '인센티브': ['투자인센티브'],
    '인프라': ['전력인프라'],
    '자전거': ['전기자전거'],
    '전기차': ['신에너지차', '전기오토바이', '전기차부품', '전기차충전기', '전기차충전소'],
    '조달': ['공공조달'],
    '지식재산권': ['지식재산권침해'],
    '취업': ['취업멘토링'],
    '침체': ['경기침체'],
    '탄소': ['탄소배출'],
    '화장품': ['천연화장품', '코스메틱', '화장품규제', '화장품원료'],
    '환경': ['환경보호'],
    '흑연': ['천연흑연'],
}
for _up, _kids in KEPT_CHILDREN.items():
    MANUAL_CHILDREN.setdefault(_up, []).extend(k for k in _kids if k not in MANUAL_CHILDREN.get(_up, []))
# 팀 지정 (2026-10-21): 상하위가 아니라 각각 따로 상위어인 쌍
MANUAL_SEPARATE = [("물", "물가"), ("발전", "발전소"), ("서비스", "서비스업"), ("소재", "신소재"), ("전기", "전기차"),
                   ("트렌드", "소비트렌드"), ("온라인", "온라인쇼핑"), ("뷰티", "화장품"), ("현대", "현대화"),
                   ("리튬", "리튬이온배터리"), ("석유", "석유화학"), ("스마트", "스마트폰"),
                   ("쇼핑", "온라인쇼핑"),   # 온라인쇼핑은 팀이 상위어로 정함 (기존 사전 유지)
                   # 팀 결정 (2026-10-22): 규칙이 새로 찾았지만 뜻이 달라 각각 상위어로 둠
                   ("바이오", "바이오매스"), ("외국인", "외국인투자"), ("그린", "그린딜"), ("스마트", "스마트그리드"),
                   ("데이터", "데이터센터"), ("전기차", "미래차")]
TEAM_NOTE = "팀 지정 (2026-10-21)"
CO_UP, CO_DOWN, CO_RATIO, CO_MIN = 0.3, 0.1, 3, 3


def read_clean(path=DATA / "news_keywords_clean.csv"):
    """(표기 → 기사 번호 집합 (kw_clean: 동의어 적용 전), 기사마다 주제어 목록 (kw_sub: 상위어로 바꾸기 전))."""
    df = pd.read_csv(path, dtype=str, keep_default_na=False, usecols=lambda c: c in ("kw_clean", "kw_sub", "kw_topic"))
    col = "kw_sub" if "kw_sub" in df.columns else "kw_topic"
    arts = defaultdict(set)
    for i, v in enumerate(df["kw_clean"]):
        for w in set(x for x in v.split("|") if x):
            arts[w].add(i)
    subs = [sorted(set(x for x in v.split("|") if x)) for v in df[col]]
    return arts, subs


def openai_embed(words, env=ENV_PATH):
    """주제어 → 길이 1로 맞춘 벡터. VECTORS에 저장한 것은 다시 쓰고 새 말만 OpenAI로 임베딩합니다."""
    have = {}
    if VECTORS.exists():
        with np.load(VECTORS, allow_pickle=False) as z:   # 닫아야 아래에서 같은 파일에 다시 저장할 수 있음 (Windows)
            if str(z["model"]) == EMBED_MODEL:
                have = dict(zip(z["words"].tolist(), np.array(z["vectors"])))
    new = [w for w in words if w not in have]
    if new:
        from dotenv import dotenv_values
        from openai import OpenAI
        key = {k.lower(): v for k, v in dotenv_values(env).items()}.get("openai_api_key", "").strip()
        if not key:
            raise ValueError(f"{env}에 OPENAI_API_KEY가 없습니다.")
        client = OpenAI(api_key=key, max_retries=8)
        for i in range(0, len(new), 500):
            got = client.embeddings.create(model=EMBED_MODEL, input=new[i:i + 500]).data
            have.update((w, np.asarray(d.embedding, dtype=np.float32)) for w, d in zip(new[i:i + 500], got))
        VECTORS.parent.mkdir(parents=True, exist_ok=True)
        np.savez(VECTORS, model=EMBED_MODEL, words=np.array(list(have)), vectors=np.stack(list(have.values())))
    V = np.stack([have[w] for w in words]).astype(np.float64)
    return V / np.linalg.norm(V, axis=1, keepdims=True)


def rule_pairs(subs, embed=openai_embed):
    """(임베딩 이웃 쌍, 동시 등장 쌍) — 각각 [(a, b)], 유사도 높은 순. embed: 말 목록 → 길이 1 벡터 (테스트에서 바꿔 넣음)."""
    counts = Counter(w for a in subs for w in a)
    words = sorted((w for w, c in counts.items() if c >= EMB_MIN), key=lambda w: (-counts[w], w))
    if len(words) < 2:
        return [], []
    V = embed(words)
    S = V @ V.T
    np.fill_diagonal(S, -1)
    top = {}
    for i in range(len(words)):
        for j in np.argsort(-S[i], kind="stable")[:EMB_TOP]:
            top[tuple(sorted((i, int(j))))] = S[i, j]
    emb = [(words[i], words[j]) for (i, j), _ in sorted(top.items(), key=lambda x: (-x[1], x[0]))[:EMB_PAIRS]]
    ix = {w: k for k, w in enumerate(words)}
    lists = [[w for w in a if w not in rd.BROAD_WORDS] for a in subs]
    dfc = Counter(w for a in lists for w in a)
    pc = Counter(p for a in lists for p in combinations([w for w in a if dfc[w] >= CO_MIN_DOC], 2))
    co = [(a, b, S[ix[a], ix[b]]) for (a, b), c in pc.items()
          if c >= CO_MIN_BOTH and c / min(dfc[a], dfc[b]) >= CO_SHARE and a in ix and b in ix]
    co = [(a, b) for a, b, _ in sorted(co, key=lambda x: (-x[2], x[0], x[1]))[:CO_PAIRS]]
    return emb, co


def source_pairs(subs, embed=openai_embed, synonyms_path=None):
    """[(a, b, 출처)] — 동의어 사전 · 임베딩 이웃 · 동시 등장의 모든 쌍 (방향 무관, 중복 제거는 나중에)."""
    rows = []
    path = Path(synonyms_path or bs.kc.SYNONYMS_PATH)
    syn = (pd.read_csv(path, dtype=str, keep_default_na=False) if path.exists()
           else pd.DataFrame(columns=["표기", "대표어", "근거"]))
    for a, b, why in zip(syn["표기"], syn["대표어"], syn["근거"]):
        # 접미사 규칙 줄의 대표어는 사슬 끝(인공지능산업 → AI)이라, 바로 위 기본 단어(인공지능)와의 쌍으로 봄
        suf = why[len("접미사("):-1] if why.startswith("접미사(") else ""
        if suf and a.endswith(suf) and a[: -len(suf)] != b:
            rows.append((a, a[: -len(suf)], "동의어 사전"))
        rows.append((a, b, "동의어 사전"))
    emb, co = rule_pairs(subs, embed)
    rows += [(a, b, "임베딩 이웃") for a, b in emb]
    rows += [(a, b, "동시 등장") for a, b in co]
    return [(a, b, src) for a, b, src in rows if a and b and a != b]


def group_pairs():
    """GROUPS로 같은 묶음에 든 쌍 (같은 뜻으로 확정한 표기들)."""
    out = set()
    for _, forms in bs.GROUPS:
        forms = [bs.kc.normalize(f) for f in forms]
        out |= {frozenset((x, y)) for x in forms for y in forms if x != y}
    return out


def relation(a, b, arts, groups):
    """(상위어, 하위어, 관계, 붙은 말, 근거). 같은 수준이면 상위어 = 기사 많은 쪽, 하위어 = 다른 쪽."""
    na, nb = len(arts.get(a, ())), len(arts.get(b, ()))
    big, small = (a, b) if na >= nb else (b, a)
    if bs.opposite(a, b):
        return big, small, "반대·짝", "", "반대말 쌍 (SYNONYM_OPPOSITES)"
    short, long_ = sorted((a, b), key=len)
    if short in long_ and short != long_:
        i = long_.find(short)
        pre, suf = long_[:i], long_[i + len(short):]
        if pre and suf:
            return short, long_, "세부 주제(앞뒤 붙음)", f"{pre}+…+{suf}", "글자 포함"
        if pre:
            return short, long_, "종류(앞에 수식어)", pre, "글자 포함 (접두사)"
        if suf in RANGE_SUFFIX:
            return short, long_, "범위(산업·시장·업)", suf, "글자 포함 (접미사 규칙)"
        if suf in SAME_SUFFIX:
            return short, long_, "같은 대상(범용 접미사)", suf, "글자 포함 (범용 접미사)"
        return short, long_, "세부 주제(뒤에 붙음)", suf, "글자 포함 (접미사)"
    if frozenset((a, b)) in groups:
        return big, small, "동의어(같은 수준)", "", "GROUPS 같은 뜻 묶음"
    both = len(arts.get(a, set()) & arts.get(b, set()))
    nbig, nsmall = max(na, nb), min(na, nb)
    if nsmall >= CO_MIN and both / nsmall >= CO_UP and both / nbig <= CO_DOWN and nbig >= CO_RATIO * nsmall:
        return big, small, "포함(동시 등장)", "", f"P(상위|하위)={both / nsmall:.2f}, P(하위|상위)={both / nbig:.2f}"
    return big, small, "같은 수준(판단 보류)", "", f"동시 등장 {both}건"


def apply_manual(pairs, arts):
    """팀 지정 반영: ① 따로 둘 쌍은 '별개(팀 지정)'로 ② 지정한 하위어는 그 상위어 아래로만 (다른 상위어 연결은 '다른 상위어로 이동')."""
    sep = {frozenset(p) for p in MANUAL_SEPARATE}
    key = [frozenset((a, b)) for a, b in zip(pairs["상위어"], pairs["하위어"])]
    pairs.loc[[k in sep for k in key], ["관계", "근거"]] = ["별개(팀 지정)", TEAM_NOTE + ": 각각 상위어"]
    for up, children in MANUAL_CHILDREN.items():
        for down in children:
            moved = (pairs["하위어"] == down) & (pairs["상위어"] != up) & pairs["관계"].isin(HIER)
            pairs.loc[moved, "근거"] = pairs.loc[moved, "관계"] + f" → {TEAM_NOTE}: '{up}' 아래로 이동"
            pairs.loc[moved, "관계"] = "다른 상위어로 이동(팀 지정)"
            both = len(arts.get(up, set()) & arts.get(down, set()))
            row = {"상위어": up, "하위어": down, "관계": "팀 지정", "붙은 말": "", "근거": TEAM_NOTE,
                   "상위어 기사수": len(arts.get(up, ())), "하위어 기사수": len(arts.get(down, ())), "함께 나온 기사": both}
            same = (pairs["상위어"] == up) & (pairs["하위어"] == down) | (pairs["상위어"] == down) & (pairs["하위어"] == up)
            if same.any():
                i = pairs.index[same][0]
                for k, v in row.items():
                    pairs.at[i, k] = v
            else:
                pairs = pd.concat([pairs, pd.DataFrame([{**row, "출처": "팀 지정"}])], ignore_index=True)
    return pairs


COLUMNS = ["상위어", "하위어", "관계", "붙은 말", "근거", "상위어 기사수", "하위어 기사수", "함께 나온 기사", "출처"]


def build(arts=None, subs=None, embed=openai_embed, synonyms_path=None):
    if arts is None or subs is None:
        arts, subs = read_clean()
    groups = group_pairs()
    seen, rows = {}, []
    for a, b, src in source_pairs(subs, embed, synonyms_path):
        key = frozenset((a, b))
        if key in seen:
            r = rows[seen[key]]
            r["출처"] = ", ".join(dict.fromkeys(r["출처"].split(", ") + [src]))
            continue
        up, down, rel, part, why = relation(a, b, arts, groups)
        if rel in HIER and (up in rd.BROAD_WORDS or len(up) == 1):
            rel, why = EXCLUDED, f"{rel} → 상위어 '{up}'이(가) 너무 넓은 말"
        both = len(arts.get(up, set()) & arts.get(down, set()))
        seen[key] = len(rows)
        rows.append({"상위어": up, "하위어": down, "관계": rel, "붙은 말": part, "근거": why,
                     "상위어 기사수": len(arts.get(up, ())), "하위어 기사수": len(arts.get(down, ())), "함께 나온 기사": both,
                     "출처": src})
    pairs = apply_manual(pd.DataFrame(rows, columns=COLUMNS), arts)
    hier = pairs[pairs["관계"].isin(HIER)]
    # 상위어의 상위어(사슬)를 따라 맨 위 상위어를 찾음 (상위어가 여럿이면 기사 많은 쪽)
    parent = {}
    for r in hier.sort_values("상위어 기사수").itertuples(index=False):
        parent[r.하위어] = r.상위어
    def root(w, seen=()):
        return w if w not in parent or w in seen else root(parent[w], seen + (w,))
    summary = []
    for up, g in hier.groupby("상위어"):
        g = g.sort_values("하위어 기사수", ascending=False)
        kinds = g["관계"].value_counts()
        summary.append({"상위어": up, "맨 위 상위어": root(up), "상위어 기사수": int(g["상위어 기사수"].iloc[0]),
                        "하위어 개수": len(g),
                        **{k: int(kinds.get(k, 0)) for k in sorted(HIER)},
                        "하위어 예시 10개": ", ".join(f"{w}({n})" for w, n in zip(g["하위어"][:10], g["하위어 기사수"][:10]))})
    summary = pd.DataFrame(summary, columns=["상위어", "맨 위 상위어", "상위어 기사수", "하위어 개수", *sorted(HIER),
                                             "하위어 예시 10개"]).sort_values(["하위어 개수", "상위어 기사수"], ascending=False)
    return pairs, summary


def write(pairs, summary, out=OUT):
    count = pairs["관계"].value_counts().rename_axis("관계").reset_index(name="쌍 수")
    info = pd.DataFrame({"항목": ["목적", "대상", "상위어별", "전체 쌍", "같은 수준", "관계 판정", "동시 등장 기준", "주의", "팀 지정"], "내용": [
        "상위어·하위어(넓은 말 ⊃ 좁은 말) 후보와 관계 판정 — build_hypernyms.py가 상하위 쌍으로 상위어 사전을 만듦",
        "동의어 사전 + 임베딩 이웃 %d쌍 + 동시 등장 %d쌍, 중복 제거 %d쌍" % (EMB_PAIRS, CO_PAIRS, len(pairs)),
        "상위어마다 하위어 개수·관계 종류별 개수·하위어 예시 10개(기사 수 많은 순, 괄호 안은 기사 수)",
        "쌍마다 상위어·하위어·관계·붙은 말·근거·기사 수·함께 나온 기사·출처",
        "상하위가 아닌 쌍 (동의어·같은 대상·반대·판단 보류·넓은 상위어 제외)",
        "① 반대말 쌍 → ② 글자 포함(짧은 말 = 상위어, 붙은 말 위치·종류로 세분) → ③ GROUPS 같은 뜻 → ④ 동시 등장 쏠림 → ⑤ 판단 보류 "
        "→ ⑥ 넓은 말·한 글자 상위어 제외 → 마지막에 팀 지정(MANUAL_CHILDREN·MANUAL_SEPARATE)이 덮어씀",
        f"P(넓은 말|좁은 말) ≥ {CO_UP}, P(좁은 말|넓은 말) ≤ {CO_DOWN}, 넓은 말 기사 {CO_RATIO}배 이상, 좁은 말 {CO_MIN}건 이상",
        "글자 포함은 뜻을 보지 않음 (현대 ⊃ 현대차처럼 고유명사도 상하위로 나옴). 동시 등장은 번역어도 하위어로 볼 수 있음 "
        "(코스메틱 → 화장품). 틀린 결과는 MANUAL_SEPARATE·MANUAL_CHILDREN에 넣어 고침",
        "2026-10-21: " + " / ".join(f"{u} ⊃ {', '.join(c)}" for u, c in MANUAL_CHILDREN.items())
        + " | 각각 상위어: " + ", ".join(f"{a}·{b}" for a, b in MANUAL_SEPARATE)]})
    wb = Workbook()
    xu.write_sheet(wb.active, info, [14, 120]); wb.active.title = "설명"
    xu.write_sheet(wb.create_sheet("관계별 개수"), count, [26, 10])
    xu.write_sheet(wb.create_sheet("상위어별"), summary, [16, 16, 11, 10, 12, 12, 12, 12, 12, 90])
    hier = pairs[pairs["관계"].isin(HIER)].sort_values(["상위어", "하위어 기사수"], ascending=[True, False])
    xu.write_sheet(wb.create_sheet("상하위 쌍"), hier, [16, 18, 20, 10, 30, 10, 10, 10, 30])
    same = (pairs[~pairs["관계"].isin(HIER)].sort_values(["관계", "상위어 기사수"], ascending=[True, False])
            .rename(columns={"상위어": "말 A (기사 많은 쪽)", "하위어": "말 B", "상위어 기사수": "말 A 기사수", "하위어 기사수": "말 B 기사수"}))
    xu.write_sheet(wb.create_sheet("같은 수준·반대·보류"), same, [16, 18, 20, 10, 30, 10, 10, 10, 30])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return count


def main():
    parser = argparse.ArgumentParser(description="상위어·하위어 후보와 관계 판정 확인표")
    parser.add_argument("--output", type=Path, default=OUT)
    args = parser.parse_args()
    pairs, summary = build()
    count = write(pairs, summary, args.output)
    print(count.to_string(index=False))
    print(f"상위어 {len(summary):,}개 · 쌍 {len(pairs):,}개 → {args.output}")


if __name__ == "__main__":
    main()
