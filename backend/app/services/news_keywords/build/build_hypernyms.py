"""상위어 사전(hypernyms.csv: 하위어 → 상위어)을 만듭니다. 키워드를 상위어로 바꿔 적재하는 데 씁니다.

입력: hierarchy_candidates.py의 상하위 판정(규칙 + 코드 안 예외 MANUAL_CHILDREN·MANUAL_SEPARATE)
규칙
- 대상 관계: 범위(산업·시장·업) · 종류(앞에 수식어) · 세부 주제(뒤에 붙음·앞뒤 붙음) · 포함(동시 등장) · 팀 지정
  (동의어·같은 대상·반대·판단 보류·별개(팀 지정)·다른 상위어로 이동은 넣지 않음)
- 동의어 사전 다음에 적용하므로 하위어·상위어 모두 동의어 대표어로 바꿔서 봄 (같아지면 뺌)
- 상위어가 둘 이상이면: 팀 지정이 있으면 그것, 없으면 기사 많은 상위어 (주제어 kw_topic 기준 기사 수)
- 사슬(진단기기 → 의료기기 → 의료)은 맨 위 상위어로 정리 (한 번만 바꾸면 되게). 고리(A → B → A)가 생기면 멈추고 알림
- 팀이 각각 상위어로 정한 쌍(MANUAL_SEPARATE)은 사슬을 거쳐서도 같은 상위어가 되면 안 됨 → 멈추고 알림
- 팀 지정 하위어를 동의어 사전이 먼저 다른 말로 바꾸면(온라인시장 → 온라인) 그 말의 상위어를 따르고,
  팀 지정과 결과가 다르면 '팀 지정 확인' 시트에 남김 (build_synonyms.SYNONYM_BLOCK에 그 쌍을 넣어야 팀 지정대로 됨)

출력: data/news_keywords/hypernyms.csv (하위어, 상위어, 바로 위 상위어, 관계, 근거, 하위어_기사수, 상위어_기사수, 다른 상위어 후보)
      data/news_keywords/review/hypernym_review.xlsx (사전 · 상위어별 하위어 목록)

실행: python backend/app/services/news_keywords/build/build_hypernyms.py
"""
import argparse
import importlib.util
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
NK = HERE.parent
spec = importlib.util.spec_from_file_location("hierarchy_candidates", HERE / "hierarchy_candidates.py")
hc = importlib.util.module_from_spec(spec)
spec.loader.exec_module(hc)
kc = hc.bs.kc   # keyword_clean (동의어 사전 읽기·표기 정규화)
OUT = kc.HYPERNYMS_PATH   # data/news_keywords/hypernyms.csv (git 제외, 이 코드로 다시 만듦)
COLUMNS = ["하위어", "상위어", "바로 위 상위어", "관계", "근거", "하위어_기사수", "상위어_기사수", "다른 상위어 후보"]


class HypernymConflict(ValueError):
    pass


def topic_counts(path=hc.DATA / "news_keywords_clean.csv"):
    """주제어(동의어 적용 후, 상위어로 바꾸기 전) → 기사 수. 상위어를 적용한 결과(kw_sub 칸)로 다시 만들어도 같은 사전이 되게."""
    df = pd.read_csv(path, dtype=str, keep_default_na=False, usecols=lambda c: c in ("kw_topic", "kw_sub"))
    col = "kw_sub" if "kw_sub" in df.columns else "kw_topic"
    return Counter(w for v in df[col] for w in set(x for x in v.split("|") if x))


def build(pairs, counts, syn):
    rep = lambda w: syn.get(kc.normalize(w), kc.normalize(w))
    hier = pairs[pairs["관계"].isin(hc.HIER)]
    separate = {frozenset((rep(a), rep(b))) for a, b in hc.MANUAL_SEPARATE}
    cands, team_moved = {}, []
    for r in hier.itertuples(index=False):
        up, down = rep(r.상위어), rep(r.하위어)
        if r.관계 == "팀 지정" and down != kc.normalize(r.하위어):
            # 동의어 사전이 하위어를 먼저 다른 말로 바꿈 → 그 말의 상위어는 다른 규칙으로 정하고, 결과만 아래에서 확인
            team_moved.append((kc.normalize(r.하위어), down, up))
            continue
        # 동의어를 적용하면 각각 상위어로 정한 쌍이 되는 연결은 넣지 않음 (뷰티산업 → 화장품은 뷰티산업 = 뷰티라서 빼기)
        if up != down and frozenset((up, down)) not in separate:
            cands.setdefault(down, {}).setdefault(up, (r.관계, r.근거))
    parent = {}
    for down, ups in cands.items():
        team = [u for u, (rel, _) in ups.items() if rel == "팀 지정"]
        parent[down] = team[0] if team else max(ups, key=lambda u: (counts.get(u, 0), u))

    def final(w, seen=()):
        """맨 위 상위어. 사슬이 돌아오면(A → B → A, A → B → C → A) 어느 쪽이 상위어인지 정할 수 없어 멈춤."""
        p = parent.get(w)
        if p is None:
            return w
        if p in seen + (w,):
            raise HypernymConflict("상위어 사슬이 돌아옴: " + " → ".join(seen + (w, p))
                                   + " (hierarchy_candidates.py의 MANUAL_SEPARATE나 MANUAL_CHILDREN으로 한쪽을 정하세요)")
        return final(p, seen + (w,))

    rows = []
    for down, up in parent.items():
        top = final(down)
        rel, why = cands[down][up]
        others = [u for u in cands[down] if u != up]
        rows.append({"하위어": down, "상위어": top, "바로 위 상위어": up, "관계": rel, "근거": why,
                     "하위어_기사수": counts.get(down, 0), "상위어_기사수": counts.get(top, 0),
                     "다른 상위어 후보": ", ".join(f"{u}({counts.get(u, 0)})" for u in others)})
    df = pd.DataFrame(rows, columns=COLUMNS).sort_values(["상위어", "하위어_기사수"], ascending=[True, False])
    to = dict(zip(df["하위어"], df["상위어"]))
    top_of = lambda w: to.get(w, w)
    bad = [f"{a} · {b}" for a, b in hc.MANUAL_SEPARATE if top_of(rep(a)) == top_of(rep(b))]
    if bad:
        raise HypernymConflict("각각 상위어로 정한 쌍이 같은 상위어로 합쳐짐: " + ", ".join(bad))
    check = pd.DataFrame([{"팀 지정 하위어": raw, "동의어 사전 결과": syn_rep, "최종 상위어": top_of(syn_rep),
                           "팀이 정한 상위어": team_up,
                           "결과": "같음 (괜찮음)" if top_of(syn_rep) == team_up else
                                   "다름 → build_synonyms.SYNONYM_BLOCK에 이 쌍을 넣어야 팀 지정대로 됨"}
                          for raw, syn_rep, team_up in team_moved],
                         columns=["팀 지정 하위어", "동의어 사전 결과", "최종 상위어", "팀이 정한 상위어", "결과"])
    return df.reset_index(drop=True), check


def write_review(df, check, path):
    by_top = (df.groupby("상위어").agg(상위어_기사수=("상위어_기사수", "first"), 하위어_개수=("하위어", "size"),
                                       하위어_기사수_합=("하위어_기사수", "sum"),
                                       하위어=("하위어", lambda s: ", ".join(s)))
              .reset_index().sort_values(["하위어_개수", "상위어_기사수"], ascending=False))
    examples = {u: ", ".join(f"{w}({n})" for w, n in zip(g["하위어"][:10], g["하위어_기사수"][:10]))
                for u, g in df.groupby("상위어")}
    by_top["하위어 예시 10개"] = by_top["상위어"].map(examples)
    by_top = by_top.drop(columns=["하위어"])
    wb = Workbook()
    hc.xu.write_sheet(wb.active, by_top, [16, 12, 11, 14, 100]); wb.active.title = "상위어별"
    hc.xu.write_sheet(wb.create_sheet("사전 (하위어 → 상위어)"), df, [18, 16, 16, 20, 40, 11, 11, 30])
    hc.xu.write_sheet(wb.create_sheet("팀 지정 확인"), check, [18, 18, 16, 16, 50])
    # 하위어가 상위어보다 기사가 많으면 대체 뒤 더 넓고 덜 쓰이는 말로 바뀜 (법인세 → 법인, 현대화 → 현대) → 사람이 확인
    bigger = df[df["하위어_기사수"] > df["상위어_기사수"]].sort_values("하위어_기사수", ascending=False)
    hc.xu.write_sheet(wb.create_sheet("확인 필요 (하위어가 더 많음)"), bigger, [18, 16, 16, 20, 40, 11, 11, 30])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return by_top


def main():
    parser = argparse.ArgumentParser(description="상위어 사전(hypernyms.csv) 만들기")
    parser.add_argument("--output", type=Path, default=OUT)
    parser.add_argument("--review", type=Path, default=hc.DATA / "review" / "hypernym_review.xlsx")
    args = parser.parse_args()
    pairs, _ = hc.build()
    df, check = build(pairs, topic_counts(), kc.load_synonyms())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(args.output, index=False, encoding="utf-8-sig")
    by_top = write_review(df, check, args.review)
    moved = df["하위어_기사수"].sum()
    print(f"상위어 사전 {len(df):,}줄 (상위어 {len(by_top):,}개, 바뀌는 주제어가 붙은 기사 연인원 {moved:,}) → {args.output}")
    print(df["관계"].value_counts().to_string())
    diff = check[check["결과"].str.startswith("다름")]
    if len(diff):
        print("팀 지정과 다른 결과 (동의어 사전이 먼저 바꿈):")
        print(diff.to_string(index=False))


if __name__ == "__main__":
    main()
