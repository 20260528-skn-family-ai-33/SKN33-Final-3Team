"""결과 확인용 엑셀 서식 (머리줄 굵게·색, 틀 고정, 필터, 칸 너비, 줄바꿈)."""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="DDEBF7")


def write_sheet(ws, df, widths, wrap_cols=()):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = HEAD
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for name in wrap_cols:
        col = list(df.columns).index(name) + 1
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
