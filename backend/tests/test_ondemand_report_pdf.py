"""Requested template integration: saved scope, comparisons, legacy formats and PDF."""
import copy
import io
import os
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from sqlalchemy import delete, select, update
from backend.tests import test_backend_api as base
from backend.app.api import reports, files
from backend.app.api.database import Store
from backend.app.api.worker import run_once
from backend.app.services.report_context import build_context
from backend.app.services.report_templates import render


class OnDemandReportTest(unittest.TestCase):
    def setUp(self):
        base.BackendApiTest.setUp(self)
        base.BackendApiTest.seed_review_articles(self)
        with self.engine.begin() as c:
            c.execute(update(self.meta.tables['article']).values(source_code='HANKYUNG'))
            c.execute(self.meta.tables['article_tag'].insert(), {'article_id':6,'country_code':'US',
                'hs_code':'330499','hs_ver':'HS2022','weight':1})

    tearDown = base.BackendApiTest.tearDown
    signup = base.BackendApiTest.signup
    post = base.BackendApiTest.post

    def preview(self, **changes):
        c = {'hs_items': [{'hs_code': '330499', 'hs_ver': 'HS2022'}], 'country_codes': ['US'],
             'period': {'start': '2025-12-01', 'end': '2025-12-31'}, 'compare': 'YOY_PERIOD', **changes}
        response = self.post('reports/preview', {'conditions': c})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def context(self, uid, conditions, sections=None):
        record = {'report_id': 42, 'report_type': 'RPT_ON_DEMAND', 'conditions_json': {
            **conditions, 'sections': sections if sections is not None else ['KPI', 'TOP_COUNTRIES', 'LOCAL_ISSUES', 'GROWTH_RISK']}}
        with self.engine.begin() as connection:
            store = Store(self.app.state.database, connection)
            document = reports.build_document(store, uid, record)
            original = copy.deepcopy(document)
            ctx = build_context(store, uid, record, document)
            self.assertEqual(document, original)
            return ctx, document

    def test_comparison_modes_match_saved_document_and_month_pairing(self):
        uid = self.signup()
        for mode in ('YOY', 'YOY_PERIOD', 'MOM', 'CUSTOM', 'NONE'):
            with self.subTest(mode=mode):
                extra = {'compare_period': {'start': '2025-10-01', 'end': '2025-11-30'}} if mode == 'CUSTOM' else {}
                ctx, doc = self.context(uid, self.preview(compare=mode, **extra)['conditions'])
                k = ctx['summary']['kpi']
                self.assertEqual(k['amount_usd'], 1230000)
                self.assertEqual(k['amount_usd'], doc['sections'][1]['kpi']['amount_usd'])
                self.assertEqual(k['import_usd'], 1230000)
                self.assertEqual(k['trade_balance_usd'], 0)
                self.assertEqual(ctx['conditions']['period']['end'], '2025-12-31')
                self.assertEqual(ctx['conditions']['compare']['monthly_paired'], mode not in ('NONE', 'CUSTOM'))
                self.assertEqual(k['prev_amount_usd'], {'YOY':1110000, 'YOY_PERIOD':1110000, 'MOM':1220000, 'CUSTOM':2430000, 'NONE':None}[mode])
                # API/DOCX/XLSX and PDF must expose the same selected comparison.
                self.assertEqual(doc['sections'][1]['kpi']['comparison_pct'], k['change_pct'])
                self.assertEqual(doc['sections'][1]['kpi']['previous_amount_usd'], k['prev_amount_usd'])
                html = render('RPT_ON_DEMAND', ctx)
                self.assertNotIn('템플릿 예시', html)
                self.assertNotIn('문장 속 수치 검증 통과', html)
                if mode in ('NONE', 'CUSTOM'): self.assertTrue(all(v is None for v in ctx['monthly']['previous']))

    def test_missing_comparison_is_not_zero_or_comparison_none(self):
        uid = self.signup(); conditions = self.preview()['conditions']
        with self.engine.begin() as c:
            c.execute(self.meta.tables['batch_run'].insert(), {'source_code':'CUSTOMS','job_name':'HS6_AGGREGATION',
                'target_table':'export_summary','period_start':date(2024,12,1),'period_end':date(2024,12,31),
                'status':'FAILED','detail_json':{}})
        ctx, _ = self.context(uid, conditions)
        self.assertEqual(ctx['conditions']['compare']['status'], 'NOT_READY')
        self.assertIsNone(ctx['summary']['kpi']['prev_amount_usd'])
        self.assertIn('비교 데이터 준비 중', render('RPT_ON_DEMAND', ctx))

    def test_zero_current_and_previous_only_countries_remain_in_other_total(self):
        uid = self.signup(); conditions = self.preview()['conditions']
        with self.engine.begin() as c:
            t = self.meta.tables['export_summary']
            c.execute(delete(t).where(t.c.trade_month == date(2025,12,1), t.c.country_code == 'US'))
        ctx, _ = self.context(uid, conditions)
        self.assertEqual(ctx['summary']['kpi']['amount_usd'], 0)
        self.assertEqual(ctx['summary']['kpi']['country_count'], 0)
        self.assertEqual(ctx['others']['prev_amount_usd'], 1110000)
        self.assertEqual(ctx['top_countries'], [])
        self.assertIn('수출 실적이 없습니다.', render('RPT_ON_DEMAND', ctx))

    def test_ranks_and_concentration_use_same_country_scope(self):
        uid = self.signup()
        with self.engine.begin() as c:
            t = self.meta.tables['export_summary']
            c.execute(update(t).where(t.c.trade_month == date(2025,12,1), t.c.country_code == 'US').values(amount_usd=3000000))
        ctx, _ = self.context(uid, self.preview(country_codes=[])['conditions'])
        k = ctx['summary']['kpi']
        self.assertEqual(k['amount_usd'], 5230000)  # TOTAL rows must not contribute.
        self.assertEqual((k['top1_name'], k['prev_top1_name']), ('미국', '중국'))
        self.assertEqual([r['rank_move'] for r in ctx['top_countries']], [1, -1])
        self.assertEqual(ctx['concentration']['top3']['current'], 100)

    def test_explicit_zero_rows_have_same_rankings_in_pdf_and_original_tables(self):
        uid = self.signup()
        with self.engine.begin() as c:
            t = self.meta.tables['export_summary']
            c.execute(update(t).where(t.c.trade_month == date(2025,12,1), t.c.country_code == 'US').values(amount_usd=0))
        ctx, doc = self.context(uid, self.preview()['conditions'])
        top = next(s for s in doc['sections'] if s['section_code'] == 'TOP_COUNTRIES')
        self.assertEqual(top['tables'][0]['rows'], ctx['top_countries'])
        self.assertEqual(ctx['summary']['kpi']['country_count'], 0)

    def test_report_metadata_keeps_actual_article_publishers(self):
        uid = self.signup()
        ctx, doc = self.context(uid, self.preview()['conditions'])
        self.assertEqual({s['code'] for s in doc['meta']['sources']}, {'CUSTOMS', 'HANKYUNG'})
        self.assertIn('한국경제', ctx['sources'][-1]['provider'])
        with self.engine.begin() as c:
            weekly = reports.build_document(Store(self.app.state.database, c), uid, {
                'report_type':'RPT_WEEKLY', 'conditions_json': {'period': {'start':'2025-12-15','end':'2025-12-21'},
                    'hs_items':[{'hs_code':'330499','hs_ver':'HS2022'}], 'country_codes':['US']}})
        self.assertEqual(weekly['meta']['sources'], [{'code':'HANKYUNG','label':'한국경제','url':None}])

    def test_multiple_items_are_summed_once_and_all_items_title_is_correct(self):
        uid = self.signup()
        with self.engine.begin() as c:
            c.execute(self.meta.tables['hs_code'].insert(), {'hs_code':'190230','hs_ver':'HS2022','depth':6,
                'name_ko':'면류','source_edition':'HS2022'})
            t = self.meta.tables['export_summary']
            original = dict(c.execute(select(t).where(t.c.trade_month == date(2025,12,1), t.c.country_code == 'US')).mappings().one())
            c.execute(t.insert(), {**original, 'hs6_code':'190230', 'amount_usd':100})
        hs = [{'hs_code':h,'hs_ver':'HS2022'} for h in ('330499','190230')]
        ctx, _ = self.context(uid, self.preview(hs_items=hs)['conditions'])
        self.assertEqual(ctx['summary']['kpi']['amount_usd'], 1230100)
        self.assertIn('기초화장품 외 1개 품목', render('RPT_ON_DEMAND', ctx))
        ctx, _ = self.context(uid, self.preview(hs_items=[])['conditions'])
        self.assertEqual(ctx['conditions']['hs_items'], [])
        self.assertIn('전체 품목 수출 시장 분석', render('RPT_ON_DEMAND', ctx))

    def test_data_matrix_matches_independent_row_sums_and_rankings(self):
        uid = self.signup()
        countries = ['US', 'CN'] + [f'C{i:02}' for i in range(12)]
        codes = ['330499', '190230']
        raw = {}
        with self.engine.begin() as c:
            c.execute(self.meta.tables['country'].insert(), [
                {'country_code':code,'name_ko':code,'country_role':'COUNTRY'} for code in countries[2:]])
            c.execute(self.meta.tables['hs_code'].insert(), {'hs_code':'190230','hs_ver':'HS2022','depth':6,
                'name_ko':'면류','source_edition':'HS2022'})
            for flow, factor in [('export_summary', 13), ('import_summary', 17)]:
                t = self.meta.tables[flow]; c.execute(delete(t)); rows = []
                for year in (2024, 2025):
                    for month in range(1, 13):
                        for hi, hs in enumerate(codes):
                            for ci, country in enumerate(countries):
                                value = Decimal(((year-2024)*12+month)*factor+hi*11+ci*7) % 19 * Decimal('123.45')
                                rows.append({'trade_month':date(year,month,1), 'hs6_code':hs,'hs_ver':'HS2022',
                                             'country_code':country,'amount_usd':value})
                raw[flow] = rows
                c.execute(t.insert(), rows)
                c.execute(t.insert(), {'trade_month':date(2025,12,1),'hs6_code':'330499','hs_ver':'HS2022',
                                      'country_code':'TOTAL','amount_usd':999999999})
        for mode in ('NONE', 'MOM', 'YOY_PERIOD', 'CUSTOM'):
            for hs_codes, selected in [([], []), (codes, countries[:3]), (codes[:1], countries[:1])]:
                with self.subTest(mode=mode, hs=hs_codes, countries=selected):
                    p = {'start':'2025-12-01' if mode == 'MOM' else '2025-10-01', 'end':'2025-12-31'}
                    extra = {'compare_period':{'start':'2025-06-01','end':'2025-07-31'}} if mode == 'CUSTOM' else {}
                    conditions = self.preview(compare=mode, period=p, country_codes=selected,
                        hs_items=[{'hs_code':h,'hs_ver':'HS2022'} for h in hs_codes], **extra)['conditions']
                    ctx, doc = self.context(uid, conditions)
                    def totals(flow, period):
                        return {code:sum((r['amount_usd'] for r in raw[flow]
                            if r['country_code']==code and (not hs_codes or r['hs6_code'] in hs_codes)
                            and period['start'] <= r['trade_month'].isoformat() <= period['end']), Decimal(0))
                            for code in (selected or countries)}
                    current = totals('export_summary', p)
                    total = sum(current.values()); k = ctx['summary']['kpi']
                    self.assertAlmostEqual(k['amount_usd'], float(total), places=6)
                    imported = sum(totals('import_summary', p).values())
                    self.assertAlmostEqual(k['import_usd'], float(imported), places=6)
                    self.assertAlmostEqual(k['trade_balance_usd'], float(total-imported), places=6)
                    self.assertAlmostEqual(sum(ctx['monthly']['current']), float(total), places=6)
                    self.assertAlmostEqual(sum(r['amount_usd'] for r in ctx['top_countries']) + (ctx['others'] or {}).get('amount_usd',0), float(total), places=6)
                    positive = sorted((code for code in current if current[code]>0), key=lambda code:(-current[code],code))
                    self.assertEqual(k['country_count'], len(positive))
                    self.assertAlmostEqual(k['top3_share_pct'], float(sum(current[code] for code in positive[:3])*100/total), places=6)
                    self.assertAlmostEqual(doc['sections'][1]['kpi']['amount_usd'], k['amount_usd'], places=6)
                    original_top = next(s for s in doc['sections'] if s['section_code']=='TOP_COUNTRIES')['tables'][0]['rows']
                    self.assertEqual([(r['country'],r['amount_usd']) for r in original_top],
                                     [(r['name'],r['amount_usd']) for r in ctx['top_countries']])
                    if mode != 'NONE':
                        prior = totals('export_summary', conditions['compare_period']); old = sum(prior.values())
                        self.assertAlmostEqual(k['prev_amount_usd'], float(old), places=6)
                        self.assertAlmostEqual(k['change_pct'], float((total-old)*100/old), places=5)
                        self.assertAlmostEqual(sum(r['prev_amount_usd'] for r in ctx['top_countries']) + (ctx['others'] or {}).get('prev_amount_usd',0), float(old), places=6)

    def test_selected_sections_request_privacy_and_metric_are_preserved(self):
        uid = self.signup(); conditions = self.preview(metric='GROWTH')['conditions']
        conditions['session_id'] = '123'
        ctx, _ = self.context(uid, conditions, ['GROWTH_RISK'])
        html = render('RPT_ON_DEMAND', ctx)
        self.assertIn('AI 분석 대화에서 요청', html)
        self.assertIn('증감률 · 현재 계산·정렬 기준: 수출액', html)
        self.assertNotIn('monthly', ctx)
        self.assertNotIn('top_countries', ctx)
        self.assertNotIn('수출 증가·감소 주요 시장', html)
        self.assertNotIn('월별 수출 추이</h2>', html)

    def test_exported_kpi_preserves_all_seven_headers_and_values(self):
        from docx import Document
        from openpyxl import load_workbook
        uid = self.signup()
        _, document = self.context(uid, self.preview(compare='MOM')['conditions'], ['KPI'])
        source = next(s for s in document['sections'] if s['section_code']=='KPI')['tables'][0]
        headers = ['수출액', '비교 기간 수출액', '비교 증감률', '전월 대비', '전년 대비', '수출 국가 수', '최대 국가 비중']
        values = [source['rows'][0][column['key']] for column in source['columns']]
        xlsx = files.xlsx_bytes([source], document['meta'], [])
        docx = files.report_docx(document, self.settings)
        with self.subTest(format='XLSX'):
            sheet = load_workbook(io.BytesIO(xlsx)).worksheets[0]
            self.assertEqual([cell.value for cell in sheet[1]], headers)
            self.assertEqual([cell.value for cell in sheet[2]], values)
        with self.subTest(format='DOCX'):
            table = next(t for t in Document(io.BytesIO(docx)).tables if t.cell(0,0).text=='수출액')
            self.assertEqual([cell.text for cell in table.rows[0].cells], headers)
            self.assertEqual([cell.text for cell in table.rows[1].cells], [files.display(v) for v in values])
        output = os.environ.get('REPORT_EXPORT_QA_DIRECTORY')
        if output:
            folder = Path(output); folder.mkdir(parents=True, exist_ok=True)
            (folder/'kpi-export.docx').write_bytes(docx)
            (folder/'kpi-export.xlsx').write_bytes(xlsx)

    @unittest.skipUnless(os.environ.get('RUN_REPORT_PDF_TESTS') == '1', 'Requires Playwright Chromium')
    def test_api_worker_pdf_docx_xlsx_and_frozen_download(self):
        from pypdf import PdfReader
        uid = self.signup()
        preview = self.preview()
        created = self.post('reports', {'preview_id':preview['preview_id'], 'confirm_conditions':True,
            'formats':['PDF','DOCX'], 'sections':['KPI','TOP_COUNTRIES','LOCAL_ISSUES']}, **{'Idempotency-Key':'ondemand-html'})
        self.assertEqual(created.status_code, 202, created.text)
        rid = created.json()['job']['report_id']
        with patch('backend.app.api.files.reportlab_pdf', side_effect=AssertionError('no fallback')):
            self.assertTrue(run_once(self.app.state.database, self.settings))
        response = self.client.get(f'/api/v1/reports/{rid}/files/pdf')
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else '')
        pdf = PdfReader(io.BytesIO(response.content))
        text = '\n'.join(page.extract_text() for page in pdf.pages)
        self.assertIn('BORDERLESSSELLERS', ''.join(text.split()))
        self.assertIn('요청형 리포트', text)
        self.assertIn('억 달러', text)
        self.assertIn('한국경제', text)
        self.assertNotIn('기사 저작권은 KOTRA', text)
        self.assertNotIn('템플릿 예시', text)
        self.assertNotIn('\ufffd', text)
        with self.engine.begin() as c:
            saved = c.scalar(select(self.meta.tables['report'].c.content_json).where(self.meta.tables['report'].c.report_id == int(rid)))
        self.assertEqual(saved['template_pdf']['report_type'], 'RPT_ON_DEMAND')
        stripped = {k:v for k,v in saved.items() if k != 'template_pdf'}
        # DOCX consumes the original sections, not template metrics or HTML.
        from docx import Document
        plain = lambda payload: '\n'.join(p.text for p in Document(io.BytesIO(payload)).paragraphs)
        self.assertEqual(plain(files.report_docx(saved, self.settings)), plain(files.report_docx(stripped, self.settings)))
        for fmt in ('docx', 'xlsx'):
            downloaded = self.client.get(f'/api/v1/reports/{rid}/files/{fmt}')
            self.assertEqual(downloaded.status_code, 200)
            if fmt == 'docx':
                headers = [cell.text for t in Document(io.BytesIO(downloaded.content)).tables for cell in t.rows[0].cells]
            else:
                from openpyxl import load_workbook
                headers = [cell.value for cell in load_workbook(io.BytesIO(downloaded.content)).worksheets[0][1]]
            for label in ('비교 기간 수출액', '비교 증감률', '최대 국가 비중'):
                self.assertIn(label, headers)
        with patch('backend.app.api.reports.build_document', side_effect=AssertionError('no rebuild')):
            self.assertEqual(self.client.get(f'/api/v1/reports/{rid}/files/pdf').content, response.content)
        other = self.signup('other@example.org')
        self.assertNotEqual(uid, other)
        self.assertEqual(self.client.get(f'/api/v1/reports/{rid}/files/pdf').status_code, 404)
        output = os.environ.get('REPORT_PDF_QA_DIRECTORY')
        if output:
            folder = Path(output); folder.mkdir(parents=True, exist_ok=True)
            (folder/'RPT_ON_DEMAND.pdf').write_bytes(response.content)
            (folder/'RPT_ON_DEMAND.html').write_text(render('RPT_ON_DEMAND', saved['template_pdf']['context']), encoding='utf-8')


if __name__ == '__main__': unittest.main()
