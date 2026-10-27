import io
import tempfile
import unittest
from unittest.mock import patch
from datetime import date, timedelta
from pathlib import Path
from fastapi.testclient import TestClient
from sqlalchemy import select, update, delete, text, inspect
from backend.tests.api_support import test_engine, reference_data
from backend.app.api.application import create_app
from backend.app.api.config import Settings
from backend.app.api.common import now
from backend.app.api.worker import run_once
from backend.app.api.files import xlsx_bytes
from backend.app.api.database import Store
from backend.app.api import jobs


class BackendApiTest(unittest.TestCase):
    def setUp(self):
        Path('tmp/backend_build').mkdir(parents=True,exist_ok=True)
        self.directory=tempfile.TemporaryDirectory(dir='tmp/backend_build',ignore_cleanup_errors=True)
        self.engine,self.meta=test_engine(); reference_data(self.engine,self.meta)
        self.settings=Settings(values={'API_COOKIE_SECURE':'false','API_REPORT_DIRECTORY':self.directory.name})
        self.app=create_app(self.settings,self.engine); self.client=TestClient(self.app)
        self.csrf=None

    def tearDown(self):
        self.client.close(); self.engine.dispose(); self.directory.cleanup()

    def signup(self,email='alice@example.org'):
        response=self.client.post('/api/v1/auth/signup',json={'name':'홍길동','email':email,'password':'Sample!123','password_confirm':'Sample!123'})
        self.assertEqual(response.status_code,201,response.text)
        self.csrf=response.json()['csrf_token']; return response.json()['user']['user_id']

    def post(self,path,body,**kwargs):
        return self.client.post('/api/v1/'+path,json=body,headers={'X-CSRF-Token':self.csrf,**kwargs})

    def test_auth_csrf_logout_and_validation(self):
        self.assertEqual(self.client.get('/api/v1/users/me').status_code,401)
        uid=self.signup()
        self.assertEqual(self.client.get('/api/v1/users/me').json()['user']['user_id'],uid)
        self.assertEqual(self.client.patch('/api/v1/users/me',json={'landing_page':'GLOBAL'}).status_code,403)
        self.assertEqual(self.post('auth/logout',{}).status_code,204)
        self.assertEqual(self.client.get('/api/v1/users/me').status_code,401)
        result=self.client.post('/api/v1/auth/signup',json={'name':'x','email':'bad','password':'secret','password_confirm':'secret'})
        self.assertEqual(result.status_code,400); self.assertNotIn('secret',result.text)
        self.assertEqual(set(result.json()),{'code','message','details','retryable','request_id'})

    def test_origin_admin_and_removed_routes(self):
        r=self.client.post('/api/v1/auth/login',json={'email':'x@example.org','password':'Sample!123'},headers={'Origin':'https://evil.invalid'})
        self.assertEqual(r.status_code,403)
        self.signup(); self.assertEqual(self.client.get('/api/v1/admin/users').status_code,403)
        self.assertEqual(self.client.get('/api/v1/admin/reports').status_code,404)
        self.assertEqual(self.client.get('/api/v1/articles?page=0').status_code,400)
        self.assertEqual(self.client.get('/api/v1/articles?page=1&page=2').status_code,400)
        self.assertEqual(self.client.get('/api/v1/notifications?unread_only=no').status_code,400)

    def test_interest_version_and_hs_depth(self):
        self.signup(); url='/api/v1/users/me/interests'
        version=self.client.get(url).json()['interests']['version']
        body={'hs_items':[{'hs_code':'330499','hs_ver':'HS2022'}],'country_codes':['US'],'expected_version':version,'landing_page':'INTERESTED'}
        r=self.client.put(url,json=body,headers={'X-CSRF-Token':self.csrf}); self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(r.json()['interests']['country_codes'],['US'])
        self.assertEqual(self.client.put(url,json=body,headers={'X-CSRF-Token':self.csrf}).status_code,409)
        body['expected_version']=r.json()['interests']['version']; body['hs_items'][0]['hs_code']='3304'
        self.assertEqual(self.client.put(url,json=body,headers={'X-CSRF-Token':self.csrf}).status_code,400)

    def test_snapshot_ownership_and_global_denominator(self):
        self.signup()
        r=self.client.get('/api/v1/trade/kpis?month=2025-12'); self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(r.json()['kpi']['amount_usd'],3460000)
        rid=r.json()['result_id']
        body={'conditions':{'hs_items':[{'hs_code':'330499','hs_ver':'HS2022'}],'country_codes':['US'],'period':{'start':'2025-12-01','end':'2025-12-31'}}}
        r=self.post('analysis/market-shares',body); self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(r.json()['denominator_usd'],3460000)
        self.signup('bob@example.org')
        self.assertEqual(self.client.get('/api/v1/results/'+rid).status_code,404)
        self.assertEqual(self.client.get('/api/v1/results/'+rid+'/file').status_code,404)

    def test_verified_period_and_no_zero_fill(self):
        self.signup()
        t=self.meta.tables['batch_run']
        with self.engine.begin() as c:
            c.execute(t.insert().values(source_code='CUSTOMS',job_name='HS6_AGGREGATION',target_table='export_summary',period_start=date(2025,12,1),period_end=date(2025,12,31),status='FAILED',detail_json={'schema_version':1}))
        r=self.client.get('/api/v1/trade/kpis?month=2025-12')
        self.assertEqual(r.status_code,422,r.text)
        self.assertEqual(self.client.get('/api/v1/data/availability').json()['trade_end_month'],'2025-11')

    def test_notification_consent_and_popup_ack_not_read(self):
        uid=self.signup()
        r=self.client.patch('/api/v1/users/me/notification-settings',json={'notifications_enabled':True},headers={'X-CSRF-Token':self.csrf})
        self.assertEqual(r.status_code,409)
        r=self.post('users/me/consents',{'terms_type':'NOTIFICATION','terms_version':'TERMS_V1','action':'AGREE','enable_notifications':True})
        self.assertEqual(r.status_code,200,r.text)
        with self.engine.begin() as c:
            c.execute(self.meta.tables['alert_event'].insert(),{'event_id':1,'event_type':'MAJOR_ISSUE','event_key':'a'*64,'payload_json':{'target':{},'article_ids':[]},'created_at':now()})
            c.execute(self.meta.tables['user_alert'].insert(),{'user_id':int(uid),'event_id':1,'is_interest_match':True,'created_at':now()})
        r=self.client.get('/api/v1/notifications/popups'); self.assertEqual(r.status_code,200,r.text)
        delivery=r.json()['groups'][0]['delivery_id']
        self.assertEqual(self.post('notifications/popups/'+delivery+'/ack',{}).status_code,204)
        self.assertEqual(self.client.get('/api/v1/notifications/unread-count').json()['count'],1)
        self.assertEqual(self.client.get('/api/v1/notifications/popups').json()['groups'],[])

    def test_report_idempotency_worker_files_and_frozen_content(self):
        self.signup()
        conditions={'period':{'start':'2025-12-01','end':'2025-12-31'},'hs_items':[{'hs_code':'330499','hs_ver':'HS2022'}]}
        preview=self.post('reports/preview',{'conditions':conditions}); self.assertEqual(preview.status_code,200,preview.text)
        body={'preview_id':preview.json()['preview_id'],'confirm_conditions':True,'sections':['KPI','TOP_COUNTRIES'],'formats':['PDF','DOCX']}
        response=self.post('reports',body,**{'Idempotency-Key':'request-1'}); self.assertEqual(response.status_code,202,response.text)
        rid=response.json()['job']['report_id']
        self.assertEqual(self.post('reports',body,**{'Idempotency-Key':'request-1'}).json()['job']['report_id'],rid)
        other={**body,'formats':['PDF']}; self.assertEqual(self.post('reports',other,**{'Idempotency-Key':'request-1'}).status_code,409)
        self.assertTrue(run_once(self.app.state.database,self.settings))
        state=self.client.get(f'/api/v1/reports/{rid}/status'); self.assertEqual(state.json()['job']['status'],'READY',state.text)
        old=self.client.get(f'/api/v1/reports/{rid}').json()
        for fmt in ('pdf','docx','xlsx'):
            file=self.client.get(f'/api/v1/reports/{rid}/files/{fmt}'); self.assertEqual(file.status_code,200,file.text[:100] if fmt=='xlsx' else '')
            self.assertTrue(file.content.startswith(b'%PDF' if fmt=='pdf' else b'PK'))
            Path('tmp/backend_build').mkdir(parents=True,exist_ok=True)
            Path(f'tmp/backend_build/report-test.{fmt}').write_bytes(file.content)
        with self.engine.begin() as c: c.execute(update(self.meta.tables['export_summary']).values(amount_usd=1))
        self.assertEqual(self.client.get(f'/api/v1/reports/{rid}').json(),old)
        with self.engine.connect() as c: self.assertEqual(len(c.execute(select(self.meta.tables['alert_event'])).all()),0)

    def test_chat_idempotency_ownership_and_worker(self):
        self.signup()
        r=self.post('chat/sessions',{}); self.assertEqual(r.status_code,201,r.text); sid=r.json()['session']['session_id']
        body={'question':'2025년 12월 전체 품목 수출 추이','client_message_id':'message-1'}
        r=self.post(f'chat/sessions/{sid}/messages',body); self.assertEqual(r.status_code,202,r.text); job=r.json()['job']['job_id']
        self.assertEqual(self.post(f'chat/sessions/{sid}/messages',body).json()['job']['job_id'],job)
        run_once(self.app.state.database,self.settings)
        r=self.client.get('/api/v1/chat/jobs/'+job); self.assertEqual(r.json()['job']['status'],'SUCCEEDED',r.text)
        self.assertEqual(len(self.client.get(f'/api/v1/chat/sessions/{sid}/messages').json()['items']),2)
        self.signup('other@example.org'); self.assertEqual(self.client.get('/api/v1/chat/jobs/'+job).status_code,404)

    def chat_answer(self, question, key):
        sid=self.post('chat/sessions',{}).json()['session']['session_id']
        response=self.post(f'chat/sessions/{sid}/messages',{'question':question,'client_message_id':key})
        self.assertEqual(response.status_code,202,response.text)
        jid=response.json()['job']['job_id']
        run_once(self.app.state.database,self.settings)
        job=self.client.get('/api/v1/chat/jobs/'+jid).json()['job']
        self.assertEqual(job['status'],'SUCCEEDED',job)
        return job['answer']

    def seed_review_articles(self):
        days=[date(2024,12,31),date(2025,1,1),date(2025,1,31),date(2025,2,1),date(2025,10,1),date(2025,12,15)]
        with self.engine.begin() as c:
            c.execute(self.meta.tables['article'].insert(),[
                {'article_id':i,'source_code':'KOTRA_MARKET','external_id':str(i),'title':'검수 기사',
                 'published_date':day,'source_url':f'https://example.org/news/{i}','body_masked':'검수 본문',
                 'category':'경제·무역','visibility_status':'VISIBLE','admin_meta_json':{}}
                for i,day in enumerate(days,1)])

    def test_clarification_resumes_without_job_table_or_duplicate_question(self):
        self.signup()
        sid=self.post('chat/sessions',{}).json()['session']['session_id']
        jid=self.post(f'chat/sessions/{sid}/messages',{'question':'2025년 12월 화장품 수출','client_message_id':'null-check'}).json()['job']['job_id']
        run_once(self.app.state.database,self.settings)
        job=self.client.get('/api/v1/chat/jobs/'+jid).json()['job']
        self.assertEqual(job['status'],'NEEDS_INPUT')
        with self.engine.begin() as c:
            c.execute(text("""CREATE TRIGGER enforce_message_result BEFORE UPDATE OF result_json ON chat_message
                WHEN NEW.result_json IS NOT NULL AND json_type(NEW.result_json) != 'object'
                BEGIN SELECT RAISE(ABORT, 'message_result_object'); END"""))
        response=self.post(f'chat/jobs/{jid}/clarifications',{
            'clarification_id':job['clarification']['clarification_id'],'choice_ids':['3304']})
        self.assertEqual(response.status_code,202,response.text)
        with self.engine.connect() as c:
            self.assertNotIn('async_job',inspect(c).get_table_names())
            t=self.meta.tables['chat_message']
            self.assertEqual(len(c.execute(select(t).where(t.c.session_id==int(sid))).all()),1)
        run_once(self.app.state.database,self.settings)
        job=self.client.get('/api/v1/chat/jobs/'+jid).json()['job']
        self.assertEqual(job['status'],'SUCCEEDED',job)
        self.assertEqual(job['answer']['conditions']['hs_items'],[{'hs_code':'3304','hs_ver':'HS2022'}])
        self.assertEqual(len(self.client.get(f'/api/v1/chat/sessions/{sid}/messages').json()['items']),2)

    def test_chat_articles_respect_period_boundaries(self):
        self.signup(); self.seed_review_articles()
        answer=self.chat_answer('2025년 1월 전체 품목 기사','article-period')
        self.assertEqual(answer['conditions']['period'],{'start':'2025-01-01','end':'2025-01-31'})
        self.assertEqual([a['published_date'] for a in answer['articles']],['2025-01-31','2025-01-01'])

    def test_chat_articles_without_trade_data(self):
        self.signup(); self.seed_review_articles()
        with self.engine.begin() as c:
            for name in ('export_summary','import_summary','batch_run'):
                c.execute(delete(self.meta.tables[name]))
        cases=[('2025년 1월 전체 품목 기사','2025-01-01',['2025-01-31','2025-01-01']),
               ('전체 품목 기사','2025-10-01',['2025-12-15','2025-10-01']),
               ('최근 2개월 전체 품목 기사','2025-11-01',['2025-12-15'])]
        for i,(question,start,days) in enumerate(cases):
            with self.subTest(question=question):
                answer=self.chat_answer(question,f'news-only-{i}')
                self.assertEqual(answer['intent'],'EXPLAIN')
                self.assertEqual(answer['conditions']['period']['start'],start)
                self.assertEqual([a['published_date'] for a in answer['articles']],days)

    def seed_quarterly_subscription(self, uid, baseline='2025-05'):
        with self.engine.begin() as c:
            for code,kind in [('QUARTERLY','REPORT_PERIOD'),('RPT_QUARTERLY','REPORT_TYPE')]:
                c.execute(self.meta.tables['service_option'].insert(),{'option_code':code,'option_type':kind,
                    'label_ko':'분기','enabled':True,'sort_order':1,'value_json':{}})
            c.execute(self.meta.tables['report_subscription'].insert(),{'user_id':int(uid),'period_code':'QUARTERLY','enabled':True,
                'filters_json':{'enabled_since':'2025-01-01T00:00:00Z','last_observed_month':baseline,
                                'hs_items':[],'countries':['US'],'formats':['PDF']}})
            for name in ('export_summary','import_summary'):
                t=self.meta.tables[name]
                c.execute(delete(t).where(t.c.trade_month>date(2025,10,1)))
                c.execute(update(self.meta.tables['batch_run']).where(self.meta.tables['batch_run'].c.target_table==name).values(period_end=date(2025,10,31)))

    def test_quarterly_reports_catch_up_once_after_later_month(self):
        from backend.app.api.subscriptions import enqueue_due
        uid=self.signup(); self.seed_quarterly_subscription(uid)
        with self.engine.begin() as c:
            store=Store(self.app.state.database,c)
            self.assertEqual(enqueue_due(store),2)
            self.assertEqual(enqueue_due(store),0)
            records=store.rows(select(store.t('report')).order_by(store.t('report').c.report_id))
            self.assertEqual([r['conditions_json']['period'] for r in records],[
                {'start':'2025-04-01','end':'2025-06-30'}, {'start':'2025-07-01','end':'2025-09-30'}])
            self.assertEqual(records[1]['conditions_json']['compare_period'],{'start':'2024-07-01','end':'2024-09-30'})
            self.assertTrue(all(jobs.report_execution(store,r)['status']=='QUEUED' for r in records))

    def test_quarterly_reports_wait_for_all_three_months(self):
        from backend.app.api.subscriptions import enqueue_due
        uid=self.signup(); self.seed_quarterly_subscription(uid)
        with self.engine.begin() as c:
            store=Store(self.app.state.database,c)
            for name in ('export_summary','import_summary'):
                t=self.meta.tables[name]
                c.execute(delete(t).where(t.c.trade_month==date(2025,8,1)))
                c.execute(self.meta.tables['batch_run'].insert(),{'source_code':'CUSTOMS','job_name':'HS6_AGGREGATION','target_table':name,
                    'period_start':date(2025,8,1),'period_end':date(2025,8,31),'status':'FAILED','detail_json':{}})
            self.assertEqual(enqueue_due(store),1)
            self.assertEqual(enqueue_due(store),0)
            for name in ('export_summary','import_summary'):
                c.execute(self.meta.tables[name].insert(),{'trade_month':date(2025,8,1),
                    'country_code':'US','hs6_code':'330499','hs_ver':'HS2022','amount_usd':100})
                c.execute(self.meta.tables['batch_run'].insert(),{'source_code':'CUSTOMS','job_name':'HS6_AGGREGATION','target_table':name,
                    'period_start':date(2025,8,1),'period_end':date(2025,8,31),'status':'SUCCESS','detail_json':{}})
            self.assertEqual(enqueue_due(store),1)
            self.assertEqual(enqueue_due(store),0)

    def test_quarterly_reports_do_not_backfill_before_subscription_baseline(self):
        from backend.app.api.subscriptions import enqueue_due
        uid=self.signup(); self.seed_quarterly_subscription(uid,baseline='2025-06')
        with self.engine.begin() as c:
            store=Store(self.app.state.database,c)
            self.assertEqual(enqueue_due(store),1)
            record=store.one(select(store.t('report')))
            self.assertEqual(record['conditions_json']['period'],{'start':'2025-07-01','end':'2025-09-30'})

    def test_xlsx_formula_and_limit(self):
        from openpyxl import load_workbook
        t={'total_rows':1,'columns':[{'key':'text','label':'텍스트'}],'rows':[{'text':'=1+1'}]}
        wb=load_workbook(io.BytesIO(xlsx_bytes([t],{},['SELECT 1'])))
        self.assertEqual(wb['결과표1']['A2'].data_type,'s'); self.assertIn('정보',wb.sheetnames)
        from backend.app.api.common import ApiError
        t['total_rows']=1001
        with self.assertRaises(ApiError): xlsx_bytes([t],{},[])

    def test_un_comparisons_separate_valuation_and_missing_observations(self):
        self.signup(); t=self.meta.tables['un_data']
        base={'country_code':'US','hs6_code':'330499','hs_ver':'HS2022','freq_code':'M'}
        with self.engine.begin() as c:
            c.execute(t.insert(),[
                {**base,'period_start':date(2024,1,1),'valuation_basis':'FOB','total_import_usd':100,'korea_import_usd':10,'korea_share_pct':10},
                {**base,'period_start':date(2025,1,1),'valuation_basis':'CIF','total_import_usd':200,'korea_import_usd':20,'korea_share_pct':10}])
        response=self.post('analysis/import-market-shares',{'hs6_code':'330499','hs_ver':'HS2022','country_codes':['US'],'freq_code':'M','period_start':'2025-01-01','period_end':'2025-02-01'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertIsNone(response.json()['items'][0]['market_yoy_pct'])
        self.assertEqual(len(response.json()['items']),1)
        self.assertTrue(any('2025-02-01' in w for w in response.json()['warnings']))

    def test_subscription_validation_and_publication_trigger(self):
        uid=self.signup(); url='/api/v1/users/me/subscriptions'
        version=self.client.get(url).json()['version']
        body={'period_codes':['MONTHLY'],'formats':['PDF'],'use_saved_interests':True,'expected_version':version}
        self.assertEqual(self.client.put(url,json=body,headers={'X-CSRF-Token':self.csrf}).status_code,409)
        interest=self.client.get('/api/v1/users/me/interests').json()['interests']
        self.client.put('/api/v1/users/me/interests',json={'hs_items':[],'country_codes':['US'],'expected_version':interest['version']},headers={'X-CSRF-Token':self.csrf})
        self.assertEqual(self.client.put(url,json={**body,'use_saved_interests':False},headers={'X-CSRF-Token':self.csrf}).status_code,400)
        self.assertEqual(self.client.put(url,json={**body,'period_codes':[]},headers={'X-CSRF-Token':self.csrf}).status_code,400)
        response=self.client.put(url,json=body,headers={'X-CSRF-Token':self.csrf}); self.assertEqual(response.status_code,200,response.text)
        from backend.app.api.subscriptions import enqueue_due
        with self.engine.begin() as c:
            store=Store(self.app.state.database,c)
            self.assertEqual(enqueue_due(store),0)  # 가입 당시 최신월은 재발행하지 않는다.
            c.execute(self.meta.tables['service_option'].insert(),{'option_code':'RPT_MONTHLY','option_type':'REPORT_TYPE','label_ko':'월간','enabled':True,'sort_order':1,'value_json':{}})
            for target in ('export_summary','import_summary'):
                c.execute(self.meta.tables[target].insert(),{'trade_month':date(2026,1,1),'hs6_code':'330499','hs_ver':'HS2022','country_code':'US','amount_usd':1500000})
                c.execute(self.meta.tables['batch_run'].insert(),{'source_code':'CUSTOMS','job_name':'HS6_AGGREGATION','target_table':target,'period_start':date(2026,1,1),'period_end':date(2026,1,31),'status':'SUCCESS','detail_json':{'schema_version':1}})
            self.assertEqual(enqueue_due(store),1)
            self.assertEqual(enqueue_due(store),0)
        self.assertFalse(self.client.get('/api/v1/users/me/notification-settings').json()['settings']['notifications_enabled'])

    def test_articles_visibility_and_admin_audit(self):
        uid=self.signup(); m=self.meta.tables
        with self.engine.begin() as c:
            c.execute(m['article'].insert(),{'article_id':1,'source_code':'KOTRA_MARKET','external_id':'test-news','title':'화장품 무역 정책',
                'published_date':date(2025,12,1),'source_url':'https://example.org/news/1','body_masked':'검증용 기사 내용','category':'통상·규제',
                'visibility_status':'VISIBLE','admin_meta_json':{'schema_version':1,'visibility':{},'country_tags_manual':False,'hs_tags_manual':False,'classification_input':None,'classification_state':{'status':'PENDING','revision':0}}})
        response=self.client.get('/api/v1/articles?tab=MAJOR_ISSUE'); self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(len(response.json()['items']),1)
        with self.engine.begin() as c: c.execute(update(m['api_user_profile']).where(m['api_user_profile'].c.user_id==int(uid)).values(role='SUPER_ADMIN'))
        detail=self.client.get('/api/v1/admin/articles/1'); self.assertEqual(detail.status_code,200,detail.text)
        response=self.client.patch('/api/v1/admin/articles/1/visibility',json={'visibility':'EXCLUDED','confirmed':True,'reason':'검수 중','expected_version':detail.json()['version']},headers={'X-CSRF-Token':self.csrf})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.client.get('/api/v1/articles/1').status_code,404)
        self.assertEqual(self.client.get('/api/v1/admin/audit-events').json()['page']['total'],1)
        self.assertEqual(self.client.get('/api/v1/admin/notification-events/999/diagnosis?user_id='+uid).status_code,404)

    def test_worker_fencing_and_duplicate_message_after_context_change(self):
        uid=self.signup(); sid=self.post('chat/sessions',{}).json()['session']['session_id']
        body={'question':'2025년 12월 전체 품목 수출 추이','client_message_id':'same'}
        job=self.post(f'chat/sessions/{sid}/messages',body).json()['job']['job_id']
        run_once(self.app.state.database,self.settings)
        self.assertEqual(self.post(f'chat/sessions/{sid}/messages',body).json()['job']['job_id'],job)
        sid2=self.post('chat/sessions',{}).json()['session']['session_id']
        self.assertEqual(self.post(f'chat/sessions/{sid2}/messages',body).status_code,409)
        from backend.app.api.common import ApiError
        with self.engine.begin() as c:
            store=Store(self.app.state.database,c)
            old=jobs.owned_job(store,int(uid),job)
            with self.assertRaises(ApiError): jobs.finish(store,{**old,'claim_token':0},{})

    def test_chat_claim_recovery_and_stale_worker_rejection(self):
        from backend.app.api import chat
        from backend.app.api.common import ApiError
        uid=self.signup(); sid=self.post('chat/sessions',{}).json()['session']['session_id']
        self.post(f'chat/sessions/{sid}/messages',{'question':'2025년 12월 전체 품목 수출','client_message_id':'recover-chat'})
        with self.engine.begin() as c: old=jobs.claim(Store(self.app.state.database,c))
        with patch('backend.app.api.jobs.now',return_value=now()+timedelta(minutes=11)):
            with self.engine.begin() as c: recovered=jobs.claim(Store(self.app.state.database,c))
        self.assertEqual(old['job_id'],recovered['job_id'])
        self.assertGreater(recovered['claim_token'],old['claim_token'])
        with self.engine.begin() as c:
            store=Store(self.app.state.database,c)
            with self.assertRaises(ApiError): jobs.finish(store,old,{})
            result=chat.execute(store,recovered); jobs.finish(store,recovered,result)
        self.assertFalse(run_once(self.app.state.database,self.settings))
        self.assertEqual(len(self.client.get(f'/api/v1/chat/sessions/{sid}/messages').json()['items']),2)

    def test_deleted_pending_chat_is_not_executed(self):
        self.signup(); sid=self.post('chat/sessions',{}).json()['session']['session_id']
        jid=self.post(f'chat/sessions/{sid}/messages',{'question':'전체 품목 수출','client_message_id':'delete-queued'}).json()['job']['job_id']
        self.assertEqual(self.client.delete('/api/v1/chat/sessions/'+sid,headers={'X-CSRF-Token':self.csrf}).status_code,204)
        self.assertFalse(run_once(self.app.state.database,self.settings))
        self.assertEqual(self.client.get('/api/v1/chat/jobs/'+jid).status_code,404)

    def queued_report(self, key):
        preview=self.post('reports/preview',{'conditions':{'period':{'start':'2025-12-01','end':'2025-12-31'}}})
        self.assertEqual(preview.status_code,200,preview.text)
        response=self.post('reports',{'preview_id':preview.json()['preview_id'],'confirm_conditions':True,
            'sections':['KPI'],'formats':['PDF']},**{'Idempotency-Key':key})
        self.assertEqual(response.status_code,202,response.text)
        return response.json()['job']['report_id']

    def test_report_retry_round_and_restart_without_job_table(self):
        from backend.app.api.database import Database
        self.signup(); rid=self.queued_report('retry-report')
        with patch('backend.app.api.worker.reports.generate',side_effect=OSError('test failure')):
            self.assertTrue(run_once(self.app.state.database,self.settings))
            self.assertTrue(run_once(self.app.state.database,self.settings))
        status=self.client.get(f'/api/v1/reports/{rid}/status').json()['job']
        self.assertEqual(status['status'],'DELAYED'); self.assertEqual(status['retry_count'],1)
        self.assertEqual(status['error']['code'],'JOB_FAILED')
        restarted=Database(self.engine)
        self.assertFalse(run_once(restarted,self.settings))
        with patch('backend.app.api.jobs.now',return_value=now()+timedelta(minutes=6)):
            self.assertTrue(run_once(restarted,self.settings))
        status=self.client.get(f'/api/v1/reports/{rid}/status').json()['job']
        self.assertEqual(status['status'],'READY'); self.assertEqual(status['retry_count'],2)
        self.assertIsNone(status['error'])
        self.assertFalse(run_once(restarted,self.settings))

    def test_report_expired_claim_recovers_after_restart(self):
        from backend.app.api.database import Database
        self.signup(); rid=self.queued_report('recover-report')
        with self.engine.begin() as c: old=jobs.claim(Store(self.app.state.database,c))
        restarted=Database(self.engine)
        with patch('backend.app.api.jobs.now',return_value=now()+timedelta(minutes=11)):
            self.assertTrue(run_once(restarted,self.settings))
        status=self.client.get(f'/api/v1/reports/{rid}/status').json()['job']
        self.assertEqual(status['status'],'READY'); self.assertEqual(status['retry_count'],1)
        with self.engine.begin() as c:
            from backend.app.api.common import ApiError
            with self.assertRaises(ApiError): jobs.finish(Store(restarted,c),old,{})

    def test_report_recovers_when_db_commit_fails_after_file_log(self):
        from backend.app.api import reports
        self.signup(); rid=self.queued_report('rollback-report')
        with self.engine.begin() as c: job=jobs.claim(Store(self.app.state.database,c))
        with self.assertRaisesRegex(RuntimeError,'simulated rollback'):
            with self.engine.begin() as c:
                store=Store(self.app.state.database,c)
                result=reports.generate(store,job,self.settings)
                jobs.finish(store,job,result)
                raise RuntimeError('simulated rollback')
        self.assertEqual(self.client.get(f'/api/v1/reports/{rid}/status').json()['job']['status'],'RUNNING')
        with patch('backend.app.api.jobs.now',return_value=now()+timedelta(minutes=11)):
            self.assertTrue(run_once(self.app.state.database,self.settings))
        self.assertEqual(self.client.get(f'/api/v1/reports/{rid}/status').json()['job']['status'],'READY')
        self.assertFalse(run_once(self.app.state.database,self.settings))
        self.assertTrue((self.app.state.database.job_logs/f'{rid}.jsonl').is_file())

    def test_older_report_is_not_starved_by_chat(self):
        self.signup(); rid=self.queued_report('older-report')
        with self.engine.begin() as c:
            t=self.meta.tables['report']
            c.execute(update(t).where(t.c.report_id==int(rid)).values(created_at=now()-timedelta(minutes=1)))
        sid=self.post('chat/sessions',{}).json()['session']['session_id']
        self.post(f'chat/sessions/{sid}/messages',{'question':'전체 품목 수출','client_message_id':'later-chat'})
        with self.engine.begin() as c:
            job=jobs.claim(Store(self.app.state.database,c))
            self.assertEqual(job['job_type'],'API_REPORT')
            self.assertEqual(str(job['job_id']),rid)

    def test_report_preview_conflicts_and_ml_gate(self):
        self.signup(); conditions={'period':{'start':'2025-12-01','end':'2025-12-31'}}
        self.assertEqual(self.post('reports/preview',{'conditions':conditions,'result_id':'bad'}).status_code,400)
        self.assertEqual(self.post('reports/preview',{'conditions':conditions,'period_preset':'LAST_1M'}).status_code,503)
        self.assertEqual(self.post('reports/preview',{'conditions':{**conditions,'compare_period':conditions['period']}}).status_code,400)
        self.assertEqual(self.post('analysis/promising-markets',{'conditions':conditions}).status_code,422)
        self.assertEqual(self.post('analysis/market-patterns',{'conditions':conditions}).status_code,422)

    def test_chat_delete_cleans_results_and_blocks_job_read(self):
        self.signup(); sid=self.post('chat/sessions',{}).json()['session']['session_id']
        jid=self.post(f'chat/sessions/{sid}/messages',{'question':'2025년 12월 전체 품목 수출 추이','client_message_id':'del-1'}).json()['job']['job_id']
        run_once(self.app.state.database,self.settings)
        result=self.client.get('/api/v1/chat/jobs/'+jid).json()['job']['answer']['tables'][0]['result_id']
        self.assertEqual(self.client.delete('/api/v1/chat/sessions/'+sid,headers={'X-CSRF-Token':self.csrf}).status_code,204)
        self.assertEqual(self.client.get('/api/v1/results/'+result).status_code,404)
        self.assertEqual(self.client.get('/api/v1/chat/jobs/'+jid).status_code,404)

    def test_saved_interests_require_confirmation(self):
        self.signup(); scope=self.client.get('/api/v1/users/me/interests').json()['interests']
        self.client.put('/api/v1/users/me/interests',json={'hs_items':[],'country_codes':['US'],'expected_version':scope['version']},headers={'X-CSRF-Token':self.csrf})
        sid=self.post('chat/sessions',{}).json()['session']['session_id']
        jid=self.post(f'chat/sessions/{sid}/messages',{'question':'수출 추이를 알려줘','client_message_id':'confirm-1'}).json()['job']['job_id']
        run_once(self.app.state.database,self.settings)
        job=self.client.get('/api/v1/chat/jobs/'+jid).json()['job']; self.assertEqual(job['status'],'NEEDS_INPUT')
        response=self.post(f'chat/jobs/{jid}/clarifications',{'clarification_id':job['clarification']['clarification_id'],'choice_ids':['APPLY']})
        self.assertEqual(response.status_code,202,response.text)
        run_once(self.app.state.database,self.settings)
        answer=self.client.get('/api/v1/chat/jobs/'+jid).json()['job']['answer']
        self.assertEqual(answer['conditions']['country_codes'],['US'])

    def test_delivery_is_idempotent_and_respects_notification_consent(self):
        uid=self.signup(); m=self.meta.tables
        from backend.app.api.event_delivery import deliver_event
        with self.engine.begin() as c:
            c.execute(m['alert_event'].insert(),{'event_id':1,'event_type':'MAJOR_ISSUE','event_key':'b'*64,'payload_json':{},'created_at':now()})
            self.assertEqual(deliver_event(Store(self.app.state.database,c),1),0)
        self.post('users/me/consents',{'terms_type':'NOTIFICATION','terms_version':'TERMS_V1','action':'AGREE','enable_notifications':True})
        with self.engine.begin() as c:
            self.assertEqual(deliver_event(Store(self.app.state.database,c),1),1)
            self.assertEqual(deliver_event(Store(self.app.state.database,c),1),0)

    def test_rankings_and_changes_use_saved_measurements(self):
        self.signup()
        response=self.client.get('/api/v1/trade/rankings?period_start=2025-12&period_end=2025-12&page_size=5')
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['items'][0]['country_code'],'CN')
        self.assertEqual(response.json()['items'][0]['one_off']['status'],'NOT_EVALUATED')
        result=self.client.get('/api/v1/results/'+response.json()['result_id']).json()['table']
        self.assertEqual(result['rows'],response.json()['items'])
        self.assertEqual(response.json()['items'][0]['rank_status'],'SAME')
        response=self.post('analysis/changes',{'month':'2025-12','compare':'YOY'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['surges'],[])
        self.assertEqual(response.json()['drops'],[])
        self.assertEqual(self.client.get('/api/v1/trade/fast-growing-products?month=2025-12').status_code,200)

    def test_future_terms_do_not_invalidate_current_consent(self):
        self.signup()
        self.assertIsNone(self.client.get('/api/v1/users/me/notification-settings').json()['settings']['terms_version'])
        self.post('users/me/consents',{'terms_type':'NOTIFICATION','terms_version':'TERMS_V1','action':'AGREE','enable_notifications':True})
        with self.engine.begin() as c:
            c.execute(self.meta.tables['service_option'].insert(),{'option_code':'TERMS_FUTURE','option_type':'NOTIFICATION_TERMS','label_ko':'미래 약관','enabled':True,'sort_order':100,'value_json':{'body':'미래 시행 약관','effective_at':'2099-01-01T00:00:00Z'}})
        settings=self.client.get('/api/v1/users/me/notification-settings').json()['settings']
        self.assertTrue(settings['consent_current'])
        self.assertEqual(settings['terms_version'],'TERMS_V1')
        self.assertEqual(self.client.get('/api/v1/terms/notification').json()['terms']['version'],'TERMS_V1')

    def test_rate_limit_and_missing_schema(self):
        for _ in range(10):
            self.assertEqual(self.client.post('/api/v1/auth/login',json={'email':'none@example.org','password':'fail'}).status_code,401)
        self.assertEqual(self.client.post('/api/v1/auth/login',json={'email':'none@example.org','password':'fail'}).status_code,429)
        self.meta.tables['api_snapshot'].drop(self.engine)
        if 'api_snapshot' in self.app.state.database.metadata.tables:
            self.app.state.database.metadata.remove(self.app.state.database.metadata.tables['api_snapshot'])
        self.assertEqual(self.client.get('/health/ready').status_code,503)


    def test_sparse_verified_trade_returns_zero_across_apis(self):
        self.signup()
        with self.engine.begin() as c:
            c.execute(self.meta.tables['article'].insert(),{'article_id':1,'source_code':'KOTRA_MARKET','external_id':'sparse-news','title':'기사',
                'published_date':date(2025,12,1),'source_url':'https://example.org/1','body_masked':'본문','category':'경제','visibility_status':'VISIBLE','admin_meta_json':{}})
            c.execute(delete(self.meta.tables['import_summary']).where(self.meta.tables['import_summary'].c.country_code=='US'))
            c.execute(delete(self.meta.tables['export_summary']).where(self.meta.tables['export_summary'].c.country_code=='US',self.meta.tables['export_summary'].c.trade_month==date(2025,12,1)))
        trend=self.client.get('/api/v1/trade/trends?period_start=2025-11&period_end=2025-12&country_codes=US')
        self.assertEqual(trend.status_code,200,trend.text)
        self.assertEqual(trend.json()['items'][-1]['export_amount_usd'],0)
        self.assertEqual(trend.json()['items'][0]['import_amount_usd'],0)
        for path in ('trade/rankings?period_start=2025-11&period_end=2025-12','trade/rankings/US/products?month=2025-11'):
            response=self.client.get('/api/v1/'+path)
            self.assertEqual(response.status_code,200,response.text)
        response=self.post('analysis/local-issues',{'conditions':{'period':{'start':'2025-12-01','end':'2025-12-31'},'country_codes':['US']}})
        self.assertEqual(response.status_code,200,response.text)
        sid=self.post('chat/sessions',{}).json()['session']['session_id']
        jid=self.post(f'chat/sessions/{sid}/messages',{'question':'2025년 12월 미국 전체 품목 수출 추이','client_message_id':'zero'}).json()['job']['job_id']
        run_once(self.app.state.database,self.settings)
        self.assertEqual(self.client.get(f'/api/v1/chat/jobs/{jid}').json()['job']['answer']['answer_status'],'ANSWERED')

    def test_wholly_empty_verified_month_and_request_cache(self):
        from backend.app.api import trade
        from sqlalchemy import event
        self.signup()
        with self.engine.begin() as c:
            for name in ('export_summary','import_summary'):
                c.execute(delete(self.meta.tables[name]).where(self.meta.tables[name].c.trade_month==date(2025,12,1)))
        response=self.client.get('/api/v1/trade/trends?period_start=2025-12&period_end=2025-12')
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['items'][0]['export_amount_usd'],0)
        queries=[]
        def record(conn,cursor,statement,params,context,many): queries.append(statement)
        with self.engine.connect() as c:
            store=Store(self.app.state.database,c)
            trade.bounds(store)
            event.listen(self.engine,'before_cursor_execute',record)
            try:
                for _ in range(6): trade.bounds(store); trade.verified_months(store,'export_summary')
                self.assertEqual(queries,[])
            finally: event.remove(self.engine,'before_cursor_execute',record)
        with self.engine.begin() as c:
            c.execute(self.meta.tables['batch_run'].insert(),{'source_code':'CUSTOMS','job_name':'HS6_AGGREGATION','target_table':'export_summary','period_start':date(2025,12,1),'period_end':date(2025,12,31),'status':'FAILED','detail_json':{}})
        self.assertNotEqual(self.client.get('/api/v1/trade/trends?period_start=2025-12&period_end=2025-12').status_code,200)

if __name__=='__main__': unittest.main()
